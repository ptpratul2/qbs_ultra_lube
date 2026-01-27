import frappe
import os
import tempfile
from bs4 import BeautifulSoup
from docx import Document
from frappe.utils import getdate


def replace_in_paragraph(paragraph, replacements):
    full_text = "".join(run.text for run in paragraph.runs)
    replaced = False

    for key, value in replacements.items():
        if key in full_text:
            full_text = full_text.replace(key, str(value))
            replaced = True

    if replaced:
        paragraph.clear()
        paragraph.add_run(full_text)


def replace_placeholders_everywhere(document, replacements):
    # normal paragraphs
    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph, replacements)

    # table cells (EXXON TEMPLATE USES THESE)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)


@frappe.whitelist()
def download_print_format_docx(doctype, name):
    # -------------------------------------------------
    # 1. Get HTML from Print Format
    # -------------------------------------------------
    html = frappe.get_print(
        doctype=doctype,
        name=name,
        print_format="Exxon Mobil COA",
        as_pdf=False
    )

    # -------------------------------------------------
    # 2. Parse HTML and extract test table
    # -------------------------------------------------
    soup = BeautifulSoup(html, "html.parser")
    test_rows = []

    for table in soup.find_all("table"):
        headers = [th.get_text(strip=True) for th in table.find_all("th")]
        if headers == [
            "Sr. No.",
            "Test Parameters",
            "UoM",
            "Test Method",
            "Test Results"
        ]:
            sr_no = 1
            for tr in table.find_all("tr")[1:]:
                tds = tr.find_all("td")
                if len(tds) == 5:
                    test_rows.append({
                        "sr": str(sr_no),
                        "parameter": tds[1].get_text(strip=True),
                        "uom": tds[2].get_text(strip=True),
                        "method": tds[3].get_text(strip=True),
                        "result": tds[4].get_text(strip=True)
                    })
                    sr_no += 1
            break

    # -------------------------------------------------
    # 3. Load DOCX TEMPLATE
    # -------------------------------------------------
    template_path = frappe.get_app_path(
        "qbs_ultra_lube", "public", "templates", "Exxon_COA.docx"
    )
    document = Document(template_path)

    # -------------------------------------------------
    # 4. Replace placeholders (THIS WAS THE ISSUE)
    # -------------------------------------------------
    doc = frappe.get_doc(doctype, name)

    mfg_date = (
        getdate(doc.date_of_mfg).strftime("%B %d, %Y")
        if doc.date_of_mfg else "-"
    )

    replacements = {
        "{{NAME}}": doc.product_name or "-",
        "{{CODE}}": doc.material_code or "-",
        "{{NO}}": doc.batch_no or "-",
        "{{DATE_MFG}}": mfg_date,
        "{{BATCH_MFG}}": getdate(doc.coa_date).strftime("%B %d, %Y") or "-",
    }


    replace_placeholders_everywhere(document, replacements)

    # -------------------------------------------------
    # 5. Fill Test Table in DOCX
    # -------------------------------------------------
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers == [
            "Sr. No.",
            "Test Parameters",
            "UoM",
            "Test Method",
            "Test Results"
        ]:
            while len(table.rows) > 1:
                table._tbl.remove(table.rows[1]._tr)

            for row in test_rows:
                cells = table.add_row().cells
                cells[0].text = row["sr"]
                cells[1].text = row["parameter"]
                cells[2].text = row["uom"]
                cells[3].text = row["method"]
                cells[4].text = row["result"]
            break

    # -------------------------------------------------
    # 6. Return DOCX
    # -------------------------------------------------
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        document.save(tmp.name)
        with open(tmp.name, "rb") as f:
            frappe.local.response.filename = f"{name}.docx"
            frappe.local.response.filecontent = f.read()
            frappe.local.response.type = "download"

    os.remove(tmp.name)
