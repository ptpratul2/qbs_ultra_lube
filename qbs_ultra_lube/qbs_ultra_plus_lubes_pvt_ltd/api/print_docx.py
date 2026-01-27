import frappe
import os
import tempfile
from bs4 import BeautifulSoup
from docx import Document
from frappe.utils import getdate, today

@frappe.whitelist()
def download_print_format_docx(doctype, name):
    # -------------------------------------------------
    # 1. Get HTML from Print Format (WORKING METHOD)
    # -------------------------------------------------
    html = frappe.get_print(
        doctype=doctype,
        name=name,
        print_format="Shell COA",
        as_pdf=False
    )

    # -------------------------------------------------
    # 2. Parse HTML and extract test table
    # -------------------------------------------------
    soup = BeautifulSoup(html, "html.parser")
    test_rows = []

    for table in soup.find_all("table"):
        headers = [th.get_text(strip=True) for th in table.find_all("th")]
        if headers == ["Test Description", "Result", "Unit", "Method"]:
            for tr in table.find_all("tr")[1:]:
                tds = tr.find_all("td")
                if len(tds) == 4:
                    test_rows.append({
                        "test": tds[0].get_text(strip=True),
                        "result": tds[1].get_text(strip=True),
                        "unit": tds[2].get_text(strip=True),
                        "method": tds[3].get_text(strip=True)
                    })
            break

    # -------------------------------------------------
    # 3. Load DOCX TEMPLATE
    # -------------------------------------------------
    template_path = frappe.get_app_path(
        "qbs_ultra_lube", "public", "templates", "shell_COA.docx"
    )

    document = Document(template_path)

    # -------------------------------------------------
    # 4. Replace placeholders
    # -------------------------------------------------
    doc = frappe.get_doc(doctype, name)

    replacements = {
        "{{MATERIAL}}": doc.material_code or "-",
        "{{MATERIAL_DESCRIPTION}}": frappe.db.get_value("Item", doc.sample_name, "item_name"),
        "{{BATCH_NUMBER}}": doc.batch_no or "-",
        "{{DATE_MANUFACTURED}}": getdate(doc.date_of_mfg).strftime("%d %b %Y").upper() or "",
        "{{DATE_TESTED}}": getdate(doc.date_of_mfg).strftime("%d %b %Y").upper() or "",
        "{{DATE}}": getdate(doc.date_of_mfg).strftime("%d %b %Y").upper() or "",
    }

    for p in document.paragraphs:
        for key, value in replacements.items():
            if key in p.text:
                p.text = p.text.replace(key, str(value))

    # -------------------------------------------------
    # 5. Fill Test Table in DOCX
    # -------------------------------------------------
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers == ["Test Description", "Result", "Unit", "Method"]:

            # remove placeholder rows
            while len(table.rows) > 1:
                table._tbl.remove(table.rows[1]._tr)

            # insert data
            for row in test_rows:
                cells = table.add_row().cells
                cells[0].text = row["test"]
                cells[1].text = row["result"]
                cells[2].text = row["unit"]
                cells[3].text = row["method"]
            break

    # -------------------------------------------------
    # 6. Return DOCX
    # -------------------------------------------------
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        document.save(tmp.name)

        with open(tmp.name, "rb") as f:
            frappe.local.response.filename = f"{name}.docx"
            frappe.local.response.filecontent = f.read()
            frappe.local.response.type = "download"

    os.remove(tmp.name)
